import os
import tempfile

import requests
from docx import Document

API_URL = os.getenv("API_URL", "http://localhost:5000")
DOCX_PATH = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "data", "examples", "examples.docx")
)
CHUNK_SIZE = 10  # Número de párrafos por chunk


def split_docx(docx_path, chunk_size):
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    total = len(paragraphs)
    chunks = [paragraphs[i : i + chunk_size] for i in range(0, total, chunk_size)]
    return chunks


def create_temp_docx(paragraphs, idx):
    temp_doc = Document()
    for p in paragraphs:
        temp_doc.add_paragraph(p)
    with tempfile.NamedTemporaryFile(delete=False, suffix=f"_part{idx+1}.docx") as temp_file:
        temp_doc.save(temp_file.name)
        return temp_file.name


def upload_docx_chunk(docx_path, idx):
    with open(docx_path, "rb") as f:
        files = {
            "file": (
                os.path.basename(docx_path),
                f,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        try:
            resp = requests.post(f"{API_URL}/documents", files=files, timeout=60)
            resp.raise_for_status()
            print(f"[OK] Chunk {idx+1} subido. ID: {resp.json().get('document_id')}")
        except Exception as e:  # pylint: disable=broad-exception-caught
            print(f"[ERROR] Fallo al subir chunk {idx+1}: {e}")


def main():
    if not os.path.exists(DOCX_PATH):
        print(f"Error: No se encuentra el archivo en {DOCX_PATH}")
        return
    chunks = split_docx(DOCX_PATH, CHUNK_SIZE)
    print(f"Total de chunks a subir: {len(chunks)}")
    for idx, chunk in enumerate(chunks):
        temp_docx = create_temp_docx(chunk, idx)
        upload_docx_chunk(temp_docx, idx)
        os.remove(temp_docx)


if __name__ == "__main__":
    main()
