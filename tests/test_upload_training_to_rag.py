"""Unit tests for the upload_training_to_rag script utilities."""

import os
from unittest.mock import MagicMock, patch
import tempfile

from docx import Document
import upload_training_to_rag as utr

# ── split_docx ────────────────────────────────────────────────────────────────


class TestSplitDocx:
    def _make_docx(self, paragraphs: list[str]) -> str:
        """Write a real DOCX to a temp file and return the path."""
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        # pylint: disable-next=consider-using-with
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        doc.save(tmp.name)
        tmp.close()
        return tmp.name

    def test_splits_into_chunks_of_correct_size(self):
        paragraphs = [f"Paragraph {i}" for i in range(25)]
        path = self._make_docx(paragraphs)
        try:
            chunks = utr.split_docx(path, chunk_size=10)
            assert len(chunks) == 3  # 10 + 10 + 5
            assert len(chunks[0]) == 10
            assert len(chunks[1]) == 10
            assert len(chunks[2]) == 5
        finally:
            os.remove(path)

    def test_empty_docx_returns_empty_list(self):
        path = self._make_docx([])
        try:
            chunks = utr.split_docx(path, chunk_size=10)
            assert not chunks
        finally:
            os.remove(path)

    def test_single_paragraph_is_one_chunk(self):
        path = self._make_docx(["Only one paragraph here."])
        try:
            chunks = utr.split_docx(path, chunk_size=10)
            assert len(chunks) == 1
            assert chunks[0] == ["Only one paragraph here."]
        finally:
            os.remove(path)


# ── create_temp_docx ──────────────────────────────────────────────────────────


class TestCreateTempDocx:
    def test_creates_valid_docx_file(self):
        paragraphs = ["Hello", "World"]
        path = utr.create_temp_docx(paragraphs, idx=0)
        try:
            assert os.path.exists(path)
            assert path.endswith(".docx")
            doc = Document(path)
            texts = [p.text for p in doc.paragraphs if p.text]
            assert "Hello" in texts
            assert "World" in texts
        finally:
            os.remove(path)

    def test_filename_contains_idx(self):
        path = utr.create_temp_docx(["text"], idx=7)
        try:
            assert "part8" in path  # idx+1
        finally:
            os.remove(path)


# ── upload_docx_chunk ─────────────────────────────────────────────────────────


class TestUploadDocxChunk:
    def test_successful_upload_prints_ok(self, capsys):
        mock_resp = MagicMock()
        mock_resp.raise_for_status = MagicMock()
        mock_resp.json.return_value = {"document_id": "abc-123"}

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(b"FAKE DOCX BYTES")
            tmp_path = f.name

        try:
            with patch("upload_training_to_rag.requests.post", return_value=mock_resp):
                utr.upload_docx_chunk(tmp_path, idx=0)
            captured = capsys.readouterr()
            assert "[OK]" in captured.out
            assert "abc-123" in captured.out
        finally:
            os.remove(tmp_path)

    def test_failed_upload_prints_error(self, capsys):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(b"FAKE DOCX BYTES")
            tmp_path = f.name

        try:
            with patch(
                "upload_training_to_rag.requests.post", side_effect=Exception("connection refused")
            ):
                utr.upload_docx_chunk(tmp_path, idx=0)
            captured = capsys.readouterr()
            assert "[ERROR]" in captured.out
        finally:
            os.remove(tmp_path)
